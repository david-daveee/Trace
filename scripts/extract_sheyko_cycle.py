"""Read the supplied workbook without modifying it; preserve the detailed set cells."""
import json
import re
import sys
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(sys.argv[1]) if len(sys.argv)>1 else Path('sheyko.xlsx')
w = openpyxl.load_workbook(SOURCE, data_only=False)
cached = openpyxl.load_workbook(SOURCE, data_only=True)
params = w['Параметры']
maxima = dict(zip(('squat', 'bench', 'deadlift'), (params[f'C{r}'].value for r in (4, 5, 6))))
names = {4:'Присед',5:'Жим лёжа',6:'Становая тяга',11:'Тяга с плинтов',12:'Тяга до колен',13:'Присед со штангой на груди',14:'Тяга на подставке'}
for r in range(17,34):
    names[r] = params[f'C{r}'].value or params[f'B{r}'].value
names[21] = 'Брусья с весом'
step = params['C8'].value
days, differences, checked, counts = [], [], 0, []
for s in w.worksheets[1:]:
    phase = s['C2'].value
    for row in range(4, s.max_row + 1):
        week_cell, day_cell, label = (s.cell(row, c).value for c in (1,2,3))
        if week_cell:
            week = int(re.fullmatch(r'(\d+) неделя',week_cell)[1])
        if day_cell:
            number = int(re.search(r'№(\d+)',day_cell)[1])
            if 'СОРЕВНОВАНИЯ' in day_cell:
                assert number == 36 and label is None
                continue  # An event with no prescribed attempts, not a fabricated workout.
            day = {'name':f'Неделя {week} · День {(number-1)%3+1}', 'phase':phase,
                   'sourceNumber':number, 'groups':[]}
            if 'ПРОХОДКА' in day_cell:
                day['note']='Проходка: в исходном плане есть подходы до 105% от заданных максимумов.'
            days.append(day)
        match = re.fullmatch(r'=Параметры!B(\d+)',str(label))
        if not match:
            continue
        exercise = names[int(match[1])]
        prescription = s.cell(row,4).value
        actual, groups = [], []
        weighted = isinstance(prescription,str)
        for col in range(4,s.max_column+1):
            value = s.cell(row+1 if weighted else row,col).value
            if value is None:
                continue
            g = {'exercise':exercise,'sets':1,'sourceSheet':s.title,'sourceRow':row}
            if weighted:
                m = re.fullmatch(r'=ROUND\(Параметры!C([456])\*([0-9.]+)/Параметры!C8,0\)\*Параметры!C8',value)
                assert m, (s.title,row,col,value)
                lift = {4:'squat',5:'bench',6:'deadlift'}[int(m[1])]
                factor = float(m[2]); reps = s.cell(row+2,col).value
                kg = float((Decimal(str(maxima[lift]))*Decimal(m[2])/Decimal(str(step))).quantize(Decimal('1'), rounding=ROUND_HALF_UP)*Decimal(str(step)))
                g.update(lift=lift,factor=factor,round=step,kg=kg,reps=reps)
                cache = cached[s.title].cell(row+1,col).value
                if cache is not None:
                    assert abs(cache-kg)<1e-7,(s.title,row,col,cache,kg)
                    checked += 1
                actual.append((round(factor*100,5),reps))
            else:
                g.update(kg=-1,reps=value)
            assert isinstance(g['reps'],int) and g['reps']>0,(s.title,row,col,g)
            if groups and all(groups[-1].get(k)==g.get(k) for k in ('exercise','reps','kg','lift','factor','round')):
                groups[-1]['sets']+=1
            else:
                groups.append(g)
        assert groups,(s.title,row)
        if weighted:
            expected=[]
            for pct,reps,sets in re.findall(r'(\d+)% (\d+)x(\d+)',prescription):
                expected.extend([(float(pct),int(reps))]*int(sets))
            if actual!=expected:
                differences.append({'sheet':s.title,'cell':f'D{row}','summary':prescription,'detailed':actual})
        day['groups'].extend(groups)
    counts.append((s.title,len([d for d in days if d['phase']==phase])))
assert [d['sourceNumber'] for d in days]==list(range(1,36))
p={'id':'builtin-sheiko-12-week','routine':'sheiko-12-week','category':'powerlifting','mine':False,'nextDay':0,
   'name':'Шейко · 12 недель к соревнованиям',
   'description':'Цикл для разрядников из файла sheyko.xlsx: недели 1–8 — подготовка, 9–12 — соревновательный этап. 35 тренировок; затем соревнования №36 без заданных попыток. Веса — проценты от твоих максимумов, округление 0,5 кг. Подсобные веса выбираются вручную. В расхождениях между краткой записью и ячейками использованы подробные ячейки подходов.',
   'source':'sheyko.xlsx · План тренировок разрядников из книги Б. И. Шейко «Пауэрлифтинг»',
   'terminalNote':'Неделя 12 · Соревнования №36. В таблице нет весов и попыток для этого дня.',
   'maxima':maxima,'step':step,'days':days}
target=ROOT/'app/src/main/assets/sheiko_competition.json'
target.write_text(json.dumps(p,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
report={'blocks':counts,'days':len(days),'sets':sum(g['sets'] for d in days for g in d['groups']),
        'cachedWeightsVerified':checked,'differences':differences}
(ROOT/'outputs').mkdir(exist_ok=True)
(ROOT/'outputs/sheiko-cycle-audit.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False))
